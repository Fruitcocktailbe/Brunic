# Brunic trechter v3 — bouwt de launch-shortlist rechtstreeks uit de ERP-Excels.
# Vervangt v1 (trechter.py, PDF-parsing) als selectiemechanisme; de familie-
# classificatie van v1 is hier overgenomen. Regels: docs/architectuur.md §Datamodel,
# packages/catalog/README.md §Domeinregels, beslissingen 14/07 (J.A.R.V.I.S doc 20).
#
# Inputs (data/):
#   erp-signalement-artikelen-2026-07-14.xlsx                 prijzenlijst (65.569 rijen)
#   erp-verkoopstatistieken-per-artikel-2026-07-08.xlsx       verkoop per artikel, 12 mnd
#                                                             (NIET de per-familie-versie van 14/07!)
#   erp-valorisatie-voorraad-2026-07-08.xlsx                  voorraad (aanwezigheidslijst)
#   erp-signalement-leveranciers-2026-07-14.xlsx              crediteuren (440+LLL)
#   wc-product-export-2026-07-06.csv (of .xlsx)               WooCommerce-export
# Oude bestandsnamen (vóór de hernoeming van 22/08) worden als fallback herkend.
#
# Outputs (data/trechter/, telkens als .xlsx vóór mensen én .csv voor de pipeline):
#   trechter-v3-shortlist.(xlsx|csv)         de launch-set, verrijkt (prijs, leverancier, maat…)
#   trechter-v3-geclassificeerd.(xlsx|csv)   alle ERP-artikelen, gededupliceerd
#   trechter-v3-woo-niet-in-erp.(xlsx|csv)   Woo-producten zonder ERP-koppeling
#   trechter-v3-prijsloos.(xlsx|csv)         shortlist-artikelen zonder prijs (vraag Sandra)
#
# Selectie = (verkocht ∪ voorraad) minus uitsluitingen. Standaard uitgesloten:
# consignatie (v2-regel, baseline 12.765) + uitgefaseerd (*** — beslist 14/07).
# Verhuur/coupon-outlet/op-bestelling worden gemarkeerd en geteld maar NIET
# standaard uitgesloten, zodat het totaal vergelijkbaar blijft met de v2-baseline;
# strenger schiften kan met --uitsluiten verhuur,coupon-outlet,op-bestelling.
#
# Gebruik:  python trechter_v3.py [--categorieen "Gordijnen & stoffen,Behang,..."]
#                                 [--uitsluiten verhuur,coupon-outlet]

import argparse
import csv
import json
import re
import sys
from collections import Counter
from pathlib import Path

try:
    import openpyxl
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("openpyxl ontbreekt: pip install openpyxl")

REPO = Path(__file__).resolve().parents[2]  # repo-root (Brunic)
DATA = REPO / "data"


def eerste_bestaande(*paden):
    """Eerste pad dat bestaat; anders het eerste (voor de foutmelding)."""
    for p in paden:
        if p.exists():
            return p
    return paden[0]


PRIJZEN = DATA / "erp-signalement-artikelen-2026-07-14.xlsx"
VERKOOP = eerste_bestaande(
    DATA / "erp-verkoopstatistieken-per-artikel-2026-07-08.xlsx",
    DATA / "erp-verkoopstatistieken-2026-07-08.xlsx")  # oude naam
VOORRAAD = DATA / "erp-valorisatie-voorraad-2026-07-08.xlsx"
LEVERANCIERS = eerste_bestaande(
    DATA / "erp-signalement-leveranciers-2026-07-14.xlsx",
    DATA / "leveranciers" / "erp-signalement-leveranciers-2026-07-14.xlsx")  # oude locatie
WOO = DATA / "wc-product-export-2026-07-06.csv"
OUT = DATA / "trechter"

# ── familie-regels (overgenomen uit v1, geverifieerd via codeboek-docs 10 & 11) ──
CONSIGNATIE = {676, 677, 745, 746, 785, 786}
VERHUUR = {598, 599}
COUPON = {190, 191, 490, 491, 492, 493, 695, 696, 697, 793, 890}
OP_BESTELLING_EXTRA = {592, 593, 594}
DE_JAGER_UITZONDERING = 895  # gewone behangserie ondanks x95

STANDAARD_UIT = {"consignatie", "uitgefaseerd", "geen-omschrijving"}


def webcategorie(fam):
    if fam is None:
        return "?"
    if fam == 505 or fam in (506, 507) or fam in (545, 546, 547) or fam == 588:
        return "Raamdecoratie"
    g = fam // 100
    if g == 1: return "Vloeren"
    if g == 2: return "Slapen & wonen"
    if g == 3: return "Verf"
    if g == 4: return "Vasttapijt"
    if g == 5: return "Toebehoren"
    if g == 6:
        if 670 <= fam <= 688: return "Atelier & maatwerk"
        return "Gordijnen & stoffen"
    if g == 7: return "Tapijten & karpetten"
    if g == 8: return "Behang"
    if g == 9: return "Slapen & wonen"
    return "?"


def status_van(omschrijving, fam):
    if not omschrijving:
        return "geen-omschrijving"
    if omschrijving.lstrip().startswith("***"):
        return "uitgefaseerd"
    if fam is None:
        return "legacy-kort-nummer"
    if fam in CONSIGNATIE: return "consignatie"
    if fam in VERHUUR: return "verhuur"
    if fam in COUPON: return "coupon-outlet"
    if (fam % 100 >= 95 and fam != DE_JAGER_UITZONDERING) or fam in OP_BESTELLING_EXTRA:
        return "op-bestelling"
    return "ok"


# ── maat-parsing (domeinregels 09/07: breedte=min, lengte=max, M-suffix=rol) ──
MAAT_PAT = re.compile(r"(?<!\d)(\d{2,4})\s*[xX×]\s*(\d{1,4})\s*(M(?![A-Za-z])|CM)?", re.IGNORECASE)


def parse_maat(omschrijving, cat):
    """→ (breedte_cm, lengte_cm, is_rol, is_mat) — ints/1, lege strings waar n.v.t."""
    m = MAAT_PAT.search(omschrijving or "")
    if not m:
        return "", "", "", ""
    a, b = int(m.group(1)), int(m.group(2))
    suffix = (m.group(3) or "").upper()
    breedte, lengte = min(a, b), max(a, b)
    # rol: expliciete meter-suffix, of extreem langwerpig (100X25M-gevallen zonder suffix)
    is_rol = suffix == "M" or (breedte > 0 and lengte / breedte >= 8)
    if is_rol:
        return "", "", 1, ""  # rolmaten verwerpen: geen breedte/lengte-metafield
    is_mat = cat == "Tapijten & karpetten" and breedte < 80
    return breedte, lengte, "", (1 if is_mat else "")


# ── Excel-hulpjes (kolomnamen kunnen per export licht verschillen → alias-zoektocht) ──
def _cel(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _norm(h):
    return re.sub(r"[^A-Z0-9]", "", _cel(h).upper())


def open_xlsx(pad, sleutel_aliases, verplicht=True):
    """Vindt de header-rij (max 15 rijen diep) en geeft (kolomnaam→index, rij-iterator, wb)."""
    if not pad.exists():
        if not verplicht:
            return None, None, None
        sys.exit(f"Ontbrekende input: {pad}\n(data/ is device-local — zet de export er eerst in.)")
    wb = openpyxl.load_workbook(pad, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rijen = ws.iter_rows(values_only=True)
    for i, rij in enumerate(rijen):
        normed = [_norm(c) for c in rij]
        if any(n in sleutel_aliases for n in normed):
            return {n: idx for idx, n in enumerate(normed) if n}, rijen, wb
        if i >= 15:
            break
    sys.exit(f"{pad.name}: geen header-rij gevonden met een van {sorted(sleutel_aliases)}")


def kolom(kolommen, pad, *aliases, verplicht=True):
    for a in aliases:
        if a in kolommen:
            return kolommen[a]
    if verplicht:
        sys.exit(f"{pad.name}: kolom {aliases} niet gevonden; aanwezig: {sorted(kolommen)}")
    return None


def veld(rij, idx):
    return _cel(rij[idx]) if idx is not None and idx < len(rij) else ""


def getal(rij, idx):
    if idx is None or idx >= len(rij) or rij[idx] is None:
        return None
    try:
        return float(rij[idx])
    except (TypeError, ValueError):
        s = str(rij[idx]).strip().replace(".", "").replace(",", ".")
        try:
            return float(s)
        except ValueError:
            return None


NUMMER_ALIASES = {"NUMMER", "ARTIKEL", "ARTIKELNUMMER", "ARTNR"}

# kolombreedtes voor de xlsx-uitvoer (rest: header-breedte)
KOLBREEDTE = {"nummer": 12, "omschrijving": 55, "leverancier_naam": 30, "webcategorie": 20,
              "naam": 45, "status": 16, "beeld": 14, "artikelnummer": 14}


def _csv_cel(v):
    if v is None:
        return ""
    if isinstance(v, float):
        return f"{v:.2f}"
    return v


def schrijf_output(basis, velden, rijen):
    """Schrijft <basis>.csv (voor de pipeline) én <basis>.xlsx (voor mensen:
    vette vaste kopregel, filterknoppen, kolombreedtes, echte getallen)."""
    with open(basis.with_suffix(".csv"), "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(velden)
        w.writerows([[_csv_cel(c) for c in r] for r in rijen])
    wb = openpyxl.Workbook(write_only=True)
    ws = wb.create_sheet(basis.stem[:31])
    for i, v in enumerate(velden, 1):
        ws.column_dimensions[get_column_letter(i)].width = KOLBREEDTE.get(v, max(11, len(v) + 2))
    ws.freeze_panes = "A2"
    vet = Font(bold=True)
    kop = []
    for v in velden:
        cel = WriteOnlyCell(ws, value=v)
        cel.font = vet
        kop.append(cel)
    ws.append(kop)
    for r in rijen:
        ws.append(["" if c is None else c for c in r])
    ws.auto_filter.ref = f"A1:{get_column_letter(len(velden))}{len(rijen) + 1}"
    wb.save(basis.with_suffix(".xlsx"))


def main():
    ap = argparse.ArgumentParser(description="Brunic trechter v3 — launch-shortlist uit de ERP-Excels")
    ap.add_argument("--categorieen", default="",
                    help='alleen deze webcategorieën, komma-gescheiden (bv. "Gordijnen & stoffen,Behang,Vloeren,Vasttapijt,Tapijten & karpetten")')
    ap.add_argument("--uitsluiten", default="",
                    help=f"extra statussen uitsluiten bovenop {sorted(STANDAARD_UIT)}: verhuur,coupon-outlet,op-bestelling,legacy-kort-nummer")
    args = ap.parse_args()

    cat_filter = {c.strip().lower() for c in args.categorieen.split(",") if c.strip()}
    uitsluiten = STANDAARD_UIT | {s.strip().lower() for s in args.uitsluiten.split(",") if s.strip()}
    OUT.mkdir(parents=True, exist_ok=True)

    waarschuwingen = []

    # 1 ── prijzenlijst = master-artikelbestand
    kols, rijen, wb = open_xlsx(PRIJZEN, NUMMER_ALIASES)
    i_num = kolom(kols, PRIJZEN, *NUMMER_ALIASES)
    i_oms = kolom(kols, PRIJZEN, "OMSCHRIJVING", "NEDERLOMSCHRIJV", "NEDERLOMSCHR")
    i_prijs = kolom(kols, PRIJZEN, "VERKOOPPRIJS", "VERKOOPPRIJSINCL", "PRIJS")
    i_rol = kolom(kols, PRIJZEN, "PERROL", verplicht=False)
    i_vt = kolom(kols, PRIJZEN, "VT", "BTW", verplicht=False)
    i_beeld = kolom(kols, PRIJZEN, "BEELD", verplicht=False)
    master, dubbels = {}, 0
    for rij in rijen:
        nummer = veld(rij, i_num)
        if not nummer.isdigit():
            continue
        if nummer in master:
            dubbels += 1
            continue
        # prijs 0 = "nog geen prijs"-sentinel (domeinregel 10/07), dus géén geldige prijs
        prijs = getal(rij, i_prijs)
        per_rol = getal(rij, i_rol)
        master[nummer] = {
            "omschrijving": veld(rij, i_oms),
            "prijs": prijs if prijs and prijs > 0 else None,
            "per_rol": per_rol if per_rol and per_rol > 0 else None,
            "vt": veld(rij, i_vt),
            "beeld": veld(rij, i_beeld),
        }
    wb.close()

    # 2 ── verkocht (12 mnd) en voorraad → de unie
    # ⚠️ vereist de PER-ARTIKEL-export (08/07, 2.924 rijen) — de 14/07-versie is
    # per FAMILIE2 geaggregeerd en onbruikbaar voor de unie.
    verkocht = {}
    kols, rijen, wb = open_xlsx(VERKOOP, NUMMER_ALIASES, verplicht=False)
    if kols is None:
        waarschuwingen.append(
            f"{VERKOOP.name} ontbreekt: 'verkocht'-kant van de unie is LEEG — de "
            "shortlist mist de verkocht-maar-niet-in-voorraad-artikelen (±1.100 in v2)")
    else:
        i_num = kolom(kols, VERKOOP, *NUMMER_ALIASES)
        i_stuks = kolom(kols, VERKOOP, "HOEVEELH", "HOEVEELHEID", "AANTAL", verplicht=False)
        i_omzet = kolom(kols, VERKOOP, "BEDRAGNET", "BEDRAG", verplicht=False)
        for rij in rijen:
            nummer = veld(rij, i_num)
            if nummer.isdigit():
                verkocht[nummer] = (getal(rij, i_stuks), getal(rij, i_omzet))
        wb.close()

    kols, rijen, wb = open_xlsx(VOORRAAD, NUMMER_ALIASES)
    i_num = kolom(kols, VOORRAAD, *NUMMER_ALIASES)
    voorraad = {veld(rij, i_num) for rij in rijen if veld(rij, i_num).isdigit()}
    wb.close()

    unie = set(verkocht) | voorraad

    # 3 ── leveranciers: crediteur 440+LLL (6-cijferig; langere nummers = geen productcrediteur)
    lev_naam = {}
    kols, rijen, wb = open_xlsx(LEVERANCIERS, NUMMER_ALIASES, verplicht=False)
    if kols is None:
        waarschuwingen.append(f"{LEVERANCIERS.name} ontbreekt: kolom leverancier_naam blijft leeg")
    else:
        i_num = kolom(kols, LEVERANCIERS, *NUMMER_ALIASES)
        i_naam = kolom(kols, LEVERANCIERS, "NAAM")
        for rij in rijen:
            nummer = veld(rij, i_num)
            if len(nummer) == 6 and nummer.startswith("440"):
                lev_naam[nummer[3:6]] = veld(rij, i_naam)
        wb.close()

    # 4 ── WooCommerce: SKU-match + producten zonder ERP-koppeling (.csv, of .xlsx-fallback)
    woo_skus, woo_rows = set(), []
    woo_xlsx = WOO.with_suffix(".xlsx")
    if WOO.exists():
        with open(WOO, encoding="utf-8-sig", newline="") as f:
            csv.field_size_limit(10 ** 8)
            lezer = csv.DictReader(f)
            veldnamen = {_norm(k): k for k in (lezer.fieldnames or [])}
            k_sku = veldnamen.get("ARTIKELNUMMER") or veldnamen.get("SKU")
            k_type = veldnamen.get("TYPE")
            k_naam = veldnamen.get("NAAM") or veldnamen.get("NAME")
            if not k_sku:
                sys.exit(f"{WOO.name}: geen Artikelnummer/SKU-kolom; aanwezig: {lezer.fieldnames}")
            for row in lezer:
                sku = (row.get(k_sku) or "").strip()
                wtype = (row.get(k_type) or "").strip() if k_type else ""
                naam = (row.get(k_naam) or "").strip() if k_naam else ""
                woo_rows.append((sku, wtype, naam))
                if sku:
                    woo_skus.add(sku)
    elif woo_xlsx.exists():
        kols, rijen, wb = open_xlsx(woo_xlsx, {"ARTIKELNUMMER", "SKU"})
        i_sku = kolom(kols, woo_xlsx, "ARTIKELNUMMER", "SKU")
        i_type = kolom(kols, woo_xlsx, "TYPE", verplicht=False)
        i_naam = kolom(kols, woo_xlsx, "NAAM", "NAME", verplicht=False)
        for rij in rijen:
            sku, wtype, naam = veld(rij, i_sku), veld(rij, i_type), veld(rij, i_naam)
            woo_rows.append((sku, wtype, naam))
            if sku:
                woo_skus.add(sku)
        wb.close()
    else:
        waarschuwingen.append(f"{WOO.name} ontbreekt: geen Woo-match, woo-niet-in-erp blijft leeg")

    # 5 ── volledige classificatie (alle 65k, gededupliceerd) + de shortlist
    def decodeer(nummer):
        fam = int(nummer[3:6]) if len(nummer) == 9 else None
        lll = nummer[:3] if len(nummer) == 9 else ""
        return lll, fam

    geclassificeerd = []
    for nummer, m in master.items():
        lll, fam = decodeer(nummer)
        geclassificeerd.append([nummer, lll, fam if fam is not None else "",
                                webcategorie(fam), status_van(m["omschrijving"], fam),
                                m["omschrijving"], m["prijs"]])
    schrijf_output(OUT / "trechter-v3-geclassificeerd",
                   ["nummer", "leverancier", "familie", "webcategorie", "status", "omschrijving", "prijs"],
                   geclassificeerd)

    shortlist, uitgesloten_telling = [], Counter()
    for nummer in sorted(unie):
        m = master.get(nummer)
        lll, fam = decodeer(nummer)
        cat = webcategorie(fam)
        if m is None:
            status, oms, prijs, per_rol, vt, beeld = "niet-in-prijzenlijst", "", None, None, "", ""
        else:
            oms, prijs, per_rol, vt, beeld = m["omschrijving"], m["prijs"], m["per_rol"], m["vt"], m["beeld"]
            status = status_van(oms, fam)
        if status in uitsluiten:
            uitgesloten_telling[status] += 1
            continue
        if cat_filter and cat.lower() not in cat_filter:
            uitgesloten_telling["categorie-filter"] += 1
            continue
        breedte, lengte, is_rol, is_mat = parse_maat(oms, cat)
        stuks, omzet = verkocht.get(nummer, (None, None))
        bron = "beide" if nummer in verkocht and nummer in voorraad else ("verkocht" if nummer in verkocht else "voorraad")
        etalage = 1 if cat == "Raamdecoratie" else ""  # beslist 14/07: alles koopbaar behalve raamdecoratie
        shortlist.append([
            nummer, lll, lev_naam.get(lll, ""), fam if fam is not None else "", cat, oms,
            bron, status, prijs, "" if prijs is not None else 1, per_rol, vt,
            etalage, breedte, lengte, is_rol, is_mat,
            1 if nummer in woo_skus else "", stuks, omzet, beeld,
        ])

    velden = ["nummer", "leverancier_lll", "leverancier_naam", "familie", "webcategorie", "omschrijving",
              "bron", "status", "prijs", "prijs_ontbreekt", "per_rol", "vt",
              "etalage", "breedte_cm", "lengte_cm", "is_rol", "is_mat",
              "woo_match", "stuks_12m", "omzet_12m", "beeld"]
    schrijf_output(OUT / "trechter-v3-shortlist", velden, shortlist)

    prijsloos = [r for r in shortlist if r[velden.index("prijs_ontbreekt")] == 1]
    schrijf_output(OUT / "trechter-v3-prijsloos", velden, prijsloos)

    # Woo-producten (geen variaties) zonder ERP-koppeling
    woo_wees = []
    for sku, wtype, naam in woo_rows:
        if wtype.lower() in ("variation", "variatie", "productvariatie"):
            continue
        if not sku:
            woo_wees.append([sku, wtype, naam, "zonder-sku"])
        elif sku not in master:
            woo_wees.append([sku, wtype, naam, "sku-niet-in-erp"])
    schrijf_output(OUT / "trechter-v3-woo-niet-in-erp", ["artikelnummer", "type", "naam", "reden"], woo_wees)
    woo_zonder = len(woo_wees)

    # 6 ── samenvatting (vergelijk met de v2-baseline: 12.765)
    met_prijs = sum(1 for r in shortlist if r[velden.index("prijs_ontbreekt")] != 1)
    per_cat = Counter(r[velden.index("webcategorie")] for r in shortlist)
    status_in = Counter(r[velden.index("status")] for r in shortlist)
    print(json.dumps({
        "waarschuwingen": waarschuwingen,
        "prijzenlijst_artikelen": len(master),
        "prijzenlijst_dubbels_genegeerd": dubbels,
        "verkocht_12m": len(verkocht),
        "voorraad": len(voorraad),
        "unie": len(unie),
        "uitgesloten": dict(uitgesloten_telling.most_common()),
        "shortlist": len(shortlist),
        "v2_baseline_ter_vergelijking": 12765,
        "status_in_shortlist": dict(status_in.most_common()),
        "prijsdekking_pct": round(100 * met_prijs / len(shortlist), 1) if shortlist else None,
        "prijsloos": len(prijsloos),
        "per_webcategorie": dict(per_cat.most_common()),
        "woo_match_in_shortlist": sum(1 for r in shortlist if r[velden.index("woo_match")] == 1),
        "woo_producten_niet_in_erp": woo_zonder,
        "leveranciers_gematcht": len(lev_naam),
    }, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    sys.exit(main())
